"""P2-6: golden snapshots for client-facing HTML/SVG + the color-key invariant.

The goldens pin what clients actually receive: the preflight report (one per
whole-file verdict, including the once-crashing REVIEW), a single proof item
page, and the template preview SVG built from the shipped example booth.
Dynamic content is controlled per the plan: `today` is INJECTED (P2-6 threads
it through make_proof's builders), inputs are canned, and comparison happens
after whitespace normalization. Regenerate deliberately with
`pytest --update-goldens` and review the diff — an unreviewed golden change is
exactly the regression these tests exist to catch.

The color-key test enforces hard invariant 3: cyan=bleed, black=trim,
magenta=safe, orange=keep-clear, green=live, red=door — on BOTH the drawing
side (preview_templates.C) and the client-facing legend (client_templates).
"""
import json
import os
import re

import pytest

import client_templates as ct
import make_proof as mp
import preview_templates as pt
import proofer

HERE = os.path.dirname(os.path.abspath(__file__))
GOLDEN_DIR = os.path.join(HERE, "golden")
EXAMPLE_SPEC = os.path.join(HERE, "..", "examples", "1_booth_spec_example.json")

# Invariant 3 — the six guide color hex/meaning pairs. NEVER edit these to make
# a test pass: a change here means client templates and production templates
# no longer speak the same color language.
KEY = {"bleed": "#00AEEF",   # cyan   - bleed
       "trim":  "#111111",   # black  - trim (finished cut)
       "safe":  "#EC008C",   # magenta- safe area
       "keep":  "#F7941E",   # orange - keep-clear
       "live":  "#39B54A",   # green  - live art area
       "door":  "#ED1C24"}   # red    - door cut + hardware


# ---------- helpers ----------
def _norm(text):
    """Whitespace-normalize for comparison: strip line edges, collapse runs,
    drop blank lines. The goldens themselves are stored verbatim."""
    lines = (re.sub(r"[ \t]+", " ", ln.strip()) for ln in text.splitlines())
    return "\n".join(ln for ln in lines if ln)


def _check_golden(request, name, text):
    path = os.path.join(GOLDEN_DIR, name)
    if request.config.getoption("--update-goldens"):
        os.makedirs(GOLDEN_DIR, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
        return
    assert os.path.exists(path), (
        f"golden file {name} missing — run `pytest --update-goldens` once and commit it")
    with open(path, encoding="utf-8") as f:
        expected = f.read()
    assert _norm(text) == _norm(expected), (
        f"client-facing output drifted from tests/golden/{name}. If the change is "
        f"intentional, regenerate with `pytest --update-goldens` and review the diff.")


# ---------- canned inputs (no dates, no filesystem state) ----------
TODAY = "January 05, 2026"

RESULTS_PASS = {
    "size": ("PASS", 'matches "full + bleed" (80.12 x 175.32 in)'),
    "color": ("PASS", "CMYK only (C=0 M=100 Y=100 K=0 seen)"),
    "resolution": ("PASS", "effective 148 ppi at build scale (target 120-150)"),
    "fonts": ("PASS", "no live fonts - text is outlined"),
    "marks": ("PASS", "no obvious printer's marks inside the file"),
    "spelling": ("PASS", "checked 42 word(s) - no obvious misspellings"),
}
RESULTS_REVIEW = dict(RESULTS_PASS,
                      size=("WARN", 'matches trim size but no bleed detected — add 1" bleed per side'),
                      spelling=("WARN", "possible misspelling: EXIBIT (may be an acronym)"))
RESULTS_FAIL = dict(RESULTS_PASS,
                    color=("FAIL", "RGB content found - convert to CMYK before print"),
                    marks=("NA", "no TrimBox - cannot measure a marks margin"))
FIXES = [{"check": "color", "text": "Convert all art to CMYK. RGB shifts on press."},
         {"check": "size", "text": 'Add 1" bleed per side and re-export.'}]
GAPS = ["page 1: 1 image XObject could not be parsed (TypeError: bad stream)"]


def _canned_res(verdict, results):
    return {"verdict": verdict, "panel": {"name": "A"}, "how": "explicit --panel",
            "results": results, "fixes": FIXES if verdict != "PASS" else []}


META = {"specs": [("Finished size (H × W)", '173.32" × 78.12"'),
                  ("File size WITH bleed", '175.32" × 80.12"'),
                  ("Material", "fabric"), ("Sides", "single")],
        "version": "2", "prepped_by": "AM", "qc_by": "TR",
        "fulfillment": "delivery", "page": 1, "pages": 1,
        "placeholders": [], "missing": []}

SPEC = {"job": {"name": "Golden Job", "job_number": "12345"},
        "settings": {"scale": 0.5, "bleed_per_side_in": 1.0, "safe_margin_in": 4.0}}


# ---------- goldens: preflight report (one per whole-file verdict) ----------
@pytest.mark.parametrize("verdict,results,fixes,gaps", [
    ("PASS", RESULTS_PASS, None, None),
    ("REVIEW", RESULTS_REVIEW, FIXES, GAPS),   # the verdict that used to KeyError (P0-7)
    ("FAIL", RESULTS_FAIL, FIXES, None),
])
def test_golden_preflight_report(request, verdict, results, fixes, gaps):
    html_out = proofer.build_report_html("wall_a_art.pdf", "A", "explicit --panel",
                                         results, verdict, fixes=fixes, gaps=gaps)
    _check_golden(request, f"report_{verdict.lower()}.html", html_out)


# ---------- golden: one proof item page (pinned today, canned everything) ----------
def test_golden_proof_item_page(request):
    body = mp._item_body("Golden Job", _canned_res("REVIEW", RESULTS_REVIEW),
                         SPEC, "", "", META, today=TODAY)
    assert TODAY in body                       # the injected date is what renders
    _check_golden(request, "proof_item_review.html", body)


# ---------- golden: template preview SVG from the shipped example booth ----------
def test_golden_preview_svg(request):
    with open(EXAMPLE_SPEC, encoding="utf-8-sig") as f:
        spec = json.load(f)
    svg, n_panels = pt.build_svg(spec)
    assert n_panels == len(spec["panels"])
    _check_golden(request, "preview_example.svg", svg)


# ---------- determinism: goldens must be byte-stable across runs ----------
def test_golden_builders_are_byte_stable():
    a = proofer.build_report_html("f.pdf", "A", "how", RESULTS_REVIEW, "REVIEW",
                                  fixes=FIXES, gaps=GAPS)
    b = proofer.build_report_html("f.pdf", "A", "how", RESULTS_REVIEW, "REVIEW",
                                  fixes=FIXES, gaps=GAPS)
    assert a == b
    a = mp._item_body("J", _canned_res("PASS", RESULTS_PASS), SPEC, "", "", META, today=TODAY)
    b = mp._item_body("J", _canned_res("PASS", RESULTS_PASS), SPEC, "", "", META, today=TODAY)
    assert a == b
    with open(EXAMPLE_SPEC, encoding="utf-8-sig") as f:
        spec = json.load(f)
    assert pt.build_svg(spec)[0] == pt.build_svg(spec)[0]


# ---------- today threading (P2-6) ----------
def test_today_threads_through_job_document():
    items = [{"res": _canned_res("PASS", RESULTS_PASS), "thumb_b64": "",
              "panel": {"name": "A", "w": 78.12, "h": 173.32, "finish": "fabric",
                        "sided": "single"},
              "specs": META["specs"], "placeholders": [], "missing": []}]
    out = mp.build_job_html("Golden Job", SPEC, items, "", dict(META), today=TODAY)
    assert out.count(TODAY) >= 2               # cover + item page both use it
    import datetime
    real = datetime.date.today().strftime("%B %d, %Y")
    if real != TODAY:
        assert real not in out


def test_log_proof_accepts_injected_today(tmp_path, monkeypatch):
    import datetime
    import openpyxl
    log = tmp_path / "proof_log.xlsx"
    monkeypatch.setenv("SEE_PROOF_LOG", str(log))
    mp.log_proof("J", "1", "A", "a.pdf", "PASS", "issued", "1", "", "", "",
                 today=datetime.date(2026, 1, 5))
    ws = openpyxl.load_workbook(str(log)).active
    assert [c.value for c in ws[2]][0] == "2026-01-05"


# ---------- invariant 3: the guide color key ----------
def test_color_key_invariant_pins_all_six_pairs():
    assert pt.C == KEY, (
        "GUIDE COLOR KEY CHANGED — invariant 3 (cyan=bleed, black=trim, magenta=safe, "
        "orange=keep-clear, green=live, red=door) must not change.")


def test_client_legend_pairs_same_hexes_with_same_meanings():
    """The client-facing legend must put each meaning next to the SAME hex the
    guides are drawn with — editing either side alone fails here."""
    legend = ct._colorkey()
    for label, hexv in [("Bleed", KEY["bleed"]), ("Trim", KEY["trim"]),
                        ("Safe", KEY["safe"]), ("Keep-clear", KEY["keep"]),
                        ("Live", KEY["live"]), ("Door", KEY["door"])]:
        m = re.search(r'border-color:(#[0-9A-Fa-f]{6})"></span>' + re.escape(label), legend)
        assert m, f"legend entry starting with {label!r} not found"
        assert m.group(1).upper() == hexv.upper(), (
            f"legend pairs {label!r} with {m.group(1)}, guides draw {hexv} — "
            f"the client key and the drawn guides disagree (invariant 3)")
